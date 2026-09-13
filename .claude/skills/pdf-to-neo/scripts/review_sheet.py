# -*- coding: utf-8 -*-
"""review_sheet.py — NEO と一緒に渡す「確認箇所シート」（xlsx）を作る。

NEO の明細コメントには人向けのメモを書かない（コグニの画面・印刷に出てしまう）。
人が確かめる点・下書きが判断した点は、このシートにまとめて NEO の隣に置く（亮平さんの指示 2026-09-13）。

シート:
  確認箇所     … 要確認（人が見て決める）→ 判断（下書きが決めた。根拠）→ 参考 の順。明細 No は NEO の明細の行番号（RecordNo。rows は見積の並びで渡す）
  手入力の行   … 工賃 # / * / @、金額 * の行（report.md の「手入力の行」と同じ）
  下書きの判断 … draft_estimate の注記すべて
  合計         … 見積書の合計欄と生成 NEO の合計

openpyxl が無い PC では同じ内容の CSV（確認箇所だけ、Excel で開ける UTF-8 BOM 付き）を書く。
make_neo.py が呼ぶ。単独でも `python review_sheet.py <estimate.json> <出力.xlsx>` で作れる（生成器を呼んで行を得る）。
"""
from __future__ import annotations

import csv
import json
import os
import sys
from typing import Optional

LEVEL_ORDER = {'要確認': 0, '判断': 1, '参考': 2}
HEAD = ('No', '重要度', '区分', 'ページ', '明細No', '見積の名称', '部品コード', '内容')


def collect(est: dict, rows: Optional[list] = None, inspect_warn: Optional[list] = None, check: Optional[dict] = None,
            audit_lines: Optional[list] = None, run_out: str = '') -> list[dict]:
    """確認箇所を集める。戻り値は dict（level, kind, page, row, name, code, text）のリスト（重要度順）"""
    out: list[dict] = []
    items = est.get('items') or []
    rows = rows or []
    aligned = len(rows) == len(items)

    def _recno(i: int):
        """見積の並びで i 行目（0 始まり）の、NEO の明細 No（RecordNo）。リサイクル置換で末尾へ動いた行も NEO 上の番号を返す（Codex 指摘）"""
        if aligned and 0 <= i < len(rows):
            try:
                return int(rows[i].get('RecordNo') or (i + 1))
            except (TypeError, ValueError):
                return i + 1
        return i + 1
    for r in est.get('_review') or []:
        e = {k: r.get(k, '') for k in ('level', 'kind', 'page', 'row', 'name', 'code', 'text')}
        if isinstance(e['row'], int) and e['row'] > 0:
            e['row'] = _recno(e['row'] - 1)   # 下書きの行番号は見積（items）の並び → NEO の明細 No に直す
        out.append(e)
    veh = est.get('vehicle') or {}
    generic = str(veh.get('generic')).strip().lower() in ('true', '1', 'yes') or veh.get('generic') is True
    n_manual = 0
    for i, r in enumerate(rows):
        it = items[i] if aligned else {}
        page = it.get('_page', '') if it else ''
        code = str(r.get('PartsCode') or '').strip()
        nm = str(it.get('name') or r.get('PartsName') or '').strip()
        try:
            std, pr = int(r.get('PartsPriceStandardOutTax') or -1), int(r.get('PartsPriceOutTax') or -1)
            q = max(1, int(r.get('PartsCount') or 1))
        except (TypeError, ValueError):
            std, pr, q = -1, -1, 1
        if code and std > 0 and pr > 0 and pr != std * q:
            out.append({'level': '要確認', 'kind': '標準価格と違う', 'page': page, 'row': _recno(i), 'name': nm, 'code': code,
                        'text': f'見積 {pr:,} 円 / 標準 {std:,} 円 × {q}（{std * q:,} 円）。部品の取り違え・価格改定・数量のどれかを確かめる'})
        if not code and (it.get('manual') or r.get('_manual')):
            n_manual += 1
            if generic:  # 汎用車種（コグニ非収録）は全行が手入力なので 1 行ずつは挙げない（下でまとめて 1 件）
                continue
            memo = str(it.get('_memo') or '').strip()
            out.append({'level': '判断', 'kind': '手入力の行', 'page': page, 'row': _recno(i), 'name': nm, 'code': '',
                        'text': 'ADDATA に無い品目として手入力した' + (f'（{memo}）' if memo else '')})
    if generic and n_manual:
        out.append({'level': '判断', 'kind': '汎用車種', 'page': '', 'row': '', 'name': str(veh.get('car_name') or ''), 'code': str(veh.get('car_code') or ''),
                    'text': f'コグニ非収録の車なので汎用車種（{veh.get("car_code") or "Z10"}）で作り、明細 {n_manual} 行はすべて手入力（部品コード・標準価格なし）'})
    for w in inspect_warn or []:
        out.append({'level': '要確認', 'kind': '突合せ', 'page': '', 'row': '', 'name': '', 'code': '', 'text': str(w)})
    for w in (check or {}).get('fail') or []:
        out.append({'level': '要確認', 'kind': '紙上検算 FAIL', 'page': '', 'row': '', 'name': '', 'code': '', 'text': str(w)})
    for w in (check or {}).get('warn') or []:
        out.append({'level': '参考', 'kind': '紙上検算', 'page': '', 'row': '', 'name': '', 'code': '', 'text': str(w)})
    for ln in audit_lines or []:
        if '★' in str(ln):
            out.append({'level': '要確認', 'kind': '装備', 'page': '', 'row': '', 'name': '', 'code': '', 'text': str(ln).strip()})
    for ln in (run_out or '').splitlines():
        if '★' in ln:
            out.append({'level': '要確認', 'kind': '検算', 'page': '', 'row': '', 'name': '', 'code': '', 'text': ln.strip()})
    tt = est.get('totals') or {}
    if tt.get('tolerance'):
        out.append({'level': '要確認', 'kind': '合計の許容差', 'page': '', 'row': '', 'name': '', 'code': '',
                    'text': f"totals.tolerance {tt.get('tolerance')} 円（{tt.get('tolerance_reason') or '理由の記載なし'}）"})
    seen = set(); uniq = []
    for e in out:
        key = (e['kind'], e['row'], e['text'])
        if key not in seen:
            seen.add(key); uniq.append(e)
    uniq.sort(key=lambda e: (LEVEL_ORDER.get(e['level'], 9), e['row'] if isinstance(e['row'], int) else 1 << 30))
    return uniq


def _manual_rows(rows: list) -> list[tuple]:
    out = []
    for r in rows or []:
        wm = str(r.get('WageByManual') or ''); pm = str(r.get('PartsPriceByManual') or '')
        if wm not in ('#', '*', '@') and pm != '*':
            continue
        tm, ts = r.get('Time'), r.get('TimeStandard')
        out.append((str(r.get('PartsCode') or '----'), str(r.get('PartsName') or '').strip(), int(r.get('PartsPriceOutTax') or 0) if (r.get('PartsPriceOutTax') or 0) > 0 else '',
                    pm, int(r.get('WageOutTax') or 0) if (r.get('WageOutTax') or 0) > 0 else '', wm, tm if (tm or -1) > 0 else '', ts if (ts or 0) > 0 else '',
                    str(r.get('_std_note') or '')))
    return out


def write(path: str, entries: list[dict], est: dict, rows: Optional[list] = None, rep: Optional[dict] = None) -> str:
    """xlsx を書く（openpyxl が無ければ同名の .csv）。書いたパスを返す"""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        path = os.path.splitext(path)[0] + '.csv'
        with open(path, 'w', encoding='utf-8-sig', newline='') as fh:
            w = csv.writer(fh)
            w.writerow(HEAD)
            for i, e in enumerate(entries, start=1):
                w.writerow([i, e['level'], e['kind'], e['page'], e['row'], e['name'], e['code'], e['text']])
        return path
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '確認箇所'
    ws.append(HEAD)
    fill = {'要確認': PatternFill('solid', fgColor='FFE0E0'), '判断': PatternFill('solid', fgColor='FFF6D5'), '参考': PatternFill('solid', fgColor='EEF3FA')}
    for i, e in enumerate(entries, start=1):
        ws.append([i, e['level'], e['kind'], e['page'], e['row'], e['name'], e['code'], e['text']])
        for c in ws[ws.max_row]:
            c.alignment = Alignment(vertical='top', wrap_text=True)
        ws.cell(ws.max_row, 2).fill = fill.get(e['level'], PatternFill())
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, wdt in zip('ABCDEFGH', (5, 8, 14, 6, 7, 30, 9, 90)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = 'A2'
    ws2 = wb.create_sheet('手入力の行')
    ws2.append(('部品コード', '部品名', '金額', '金額の印', '工賃', '工賃の印', '指数', '標準指数', '説明'))
    for t in _manual_rows(rows or []):
        ws2.append(t)
    for col, wdt in zip('ABCDEFGHI', (9, 30, 10, 8, 10, 8, 7, 8, 50)):
        ws2.column_dimensions[col].width = wdt
    ws3 = wb.create_sheet('下書きの判断')
    ws3.append(('No', '内容'))
    for i, n in enumerate(est.get('_draft_notes') or [], start=1):
        ws3.append((i, str(n)))
    ws3.column_dimensions['B'].width = 140
    ws4 = wb.create_sheet('合計')
    ws4.append(('項目', '見積書', '生成 NEO'))
    pt = est.get('totals') or {}; t = (rep or {}).get('totals') or {}
    for label, k_pt, k_t in (('部品計', 'parts', 'parts'), ('工賃計', 'wage', 'wage'), ('塗装計（材料込）', 'paint_total', 'paint'), ('材料代', 'material', 'paint_material'),
                             ('内板骨格', 'frame', 'frame'), ('費用部品', 'expense_parts', 'expense_parts'), ('費用工賃', 'expense_wage', 'expense_wage'),
                             ('課税小計', 'taxable', 'subtotal'), ('消費税', 'tax', 'tax'), ('合計（税込）', 'total', 'total')):
        a = pt.get(k_pt); b = t.get(k_t)
        if a is None and not b:
            continue
        ws4.append((label, a, b))
    ws4.column_dimensions['A'].width = 18
    tmp = f'{path}.{os.getpid()}.tmp.xlsx'
    wb.save(tmp)
    os.replace(tmp, path)
    return path


def main(argv: list) -> int:
    if len(argv) < 2:
        print(__doc__ or '')
        print('使い方: python review_sheet.py <estimate.json> <出力.xlsx>')
        return 1
    here = os.path.dirname(os.path.realpath(__file__))
    sys.path.insert(0, here)
    import skill_env
    skill_env.apply()
    sys.path.insert(0, os.path.join(skill_env.FILES, 'claude_neo_pipeline'))
    from estimate_to_neo import NeoBuilder
    est = json.load(open(argv[0], encoding='utf-8-sig'))
    import run_case
    _n, rep = NeoBuilder().build(est, est['vehicle'], hints=est.get('hints'), labor_rate=est.get('labor_rate'), est_date=est.get('est_date'), insurance=est.get('insurance'))
    rows = run_case._rows_in_source_order(rep['rows'])  # collect は見積の並びの行を受け取る（リサイクル置換で末尾へ動いた行を戻す。make_neo と同じ）
    print(write(argv[1], collect(est, rows), est, rows, rep))
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
