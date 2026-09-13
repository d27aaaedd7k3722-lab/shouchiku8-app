# -*- coding: utf-8 -*-
"""確認箇所シート（review_sheet.py）の単体テスト（ADDATA 不要）
    PYTHONIOENCODING=utf-8 python .claude/skills/pdf-to-neo/scripts/tests/test_review_sheet.py
"""
from __future__ import annotations

import os
import sys
import tempfile

HERE = os.path.dirname(os.path.realpath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
import review_sheet as rs  # noqa: E402

EST = {
    'items': [{'name': 'ｸﾘｯﾌﾟ', 'code': '0170', 'qty': 13, 'price': 1300, '_page': 2},
              {'name': 'ﾌﾛﾝｶﾞｽ', 'manual': True, 'price': 50000, '_memo': '冷媒ガス。ADDATA の部品ではない', '_page': 4},
              {'name': 'ﾅｯﾄ', 'code': '0171', 'qty': 1, 'price': 140, '_page': 2}],
    '_review': [{'level': '判断', 'kind': '数量', 'page': 2, 'row': 1, 'name': 'ｸﾘｯﾌﾟ', 'code': '0170', 'text': '数量 1 → 13（金額 1,300 ÷ 標準単価 100）'},
                {'level': '要確認', 'kind': '転記メモ', 'page': 4, 'row': 2, 'name': 'ﾌﾛﾝｶﾞｽ', 'code': '', 'text': '工場の書き間違いかもしれない'}],
    '_draft_notes': ['ﾅｯﾄ: 名称だけで決めた'],
    'totals': {'parts': 51440, 'total': 56584},
}
ROWS = [{'PartsCode': '0170', 'PartsName': 'ｸﾘﾂﾌﾟ', 'PartsPriceStandardOutTax': 100, 'PartsPriceOutTax': 1300, 'PartsCount': 13},
        {'PartsCode': '', 'PartsName': 'ﾌﾛﾝｶﾞｽ', 'PartsPriceOutTax': 50000, 'PartsPriceByManual': '*', '_manual': True},
        {'PartsCode': '0171', 'PartsName': 'ﾅﾂﾄ', 'PartsPriceStandardOutTax': 90, 'PartsPriceOutTax': 140, 'PartsCount': 1, 'PartsPriceByManual': '*'}]


def test_collect_orders_by_level_and_keeps_rows():
    """要確認 → 判断 → 参考 の順。数量を直した行（1,300 = 100 × 13）は標準価格と違う行にしない。手入力の行は転記メモを理由に添える"""
    es = rs.collect(EST, ROWS, inspect_warn=['塗装 ★テスト'], check={'warn': ['WARN テスト']}, run_out='  ★ 検算の注意\n普通の行')
    levels = [e['level'] for e in es]
    assert levels == sorted(levels, key=lambda x: rs.LEVEL_ORDER[x]), levels
    kinds = [(e['kind'], e['row']) for e in es]
    assert ('標準価格と違う', 3) in kinds and ('標準価格と違う', 1) not in kinds, kinds
    man = [e for e in es if e['kind'] == '手入力の行']
    assert man and '冷媒ガス' in man[0]['text'] and man[0]['page'] == 4, man
    assert any(e['kind'] == '突合せ' for e in es) and any(e['kind'] == '検算' for e in es) and any(e['kind'] == '紙上検算' for e in es), kinds


def test_row_numbers_are_neo_record_numbers():
    """明細 No は NEO の RecordNo（リサイクル置換で末尾へ動いた行も NEO 上の番号）"""
    rows = [dict(r) for r in ROWS]
    rows[0]['RecordNo'] = 3; rows[1]['RecordNo'] = 1; rows[2]['RecordNo'] = 2  # 見積 1 行目が NEO では 3 行目
    es = rs.collect(EST, rows)
    q = [e for e in es if e['kind'] == '数量']
    assert q and q[0]['row'] == 3, q
    assert [e['row'] for e in es if e['kind'] == '標準価格と違う'] == [2], es


def test_write_xlsx_or_csv():
    d = tempfile.mkdtemp(prefix='review_')
    p = rs.write(os.path.join(d, 'x_確認箇所.xlsx'), rs.collect(EST, ROWS), EST, ROWS, {'totals': {'parts': 51440, 'total': 56584}})
    assert os.path.exists(p), p
    if p.endswith('.xlsx'):
        import openpyxl
        wb = openpyxl.load_workbook(p)
        assert wb.sheetnames == ['確認箇所', '手入力の行', '下書きの判断', '合計'], wb.sheetnames
        ws = wb['確認箇所']
        assert [c.value for c in ws[1]] == list(rs.HEAD)
        assert ws.max_row >= 4, ws.max_row
        assert wb['手入力の行'].max_row == 3, wb['手入力の行'].max_row  # 見出し + 金額 * の 2 行
    else:
        assert open(p, encoding='utf-8-sig').read().startswith('No,重要度'), p


if __name__ == '__main__':
    fails = 0
    for name, fn in sorted(globals().items()):
        if name.startswith('test_') and callable(fn):
            try:
                fn()
                print('ok  ', name)
            except AssertionError as e:
                fails += 1
                print('FAIL', name, e)
    print('review_sheet tests:', 'all ok' if not fails else f'{fails} failed')
    sys.exit(1 if fails else 0)
